#!/usr/bin/env python3
"""Make a symbolgen pins.csv from an NXP IOMUX spreadsheet (S32K3 and similar).

NXP ships the ball/pin map and pin muxing as an .xlsx attached to the device's Reference
Manual (extract it with `pdfdetach -list RM.pdf`, then `pdfdetach -save N -o out.xlsx RM.pdf`).

Usage:
  python3 scripts/pintables/nxp_iomux.py IOMUX.xlsx --device S32K344 --package 257bga \
      --type V11=power_out --type V25=power_out -o parts/<MPN>/pins.csv

Reads sheets <device>_Pinout, <device>_IO Signal Table, <device>_Supplies (columns named
<device>_<package>). Output units: "Power" (supplies on top, VSS*/VREFL at the bottom, power
outputs right, oscillator pins left), then one unit per GPIO port (bits 0-15 left, 16-31 right).

Pin rules:
  - Port pins are bidirectional; a pin whose reset default is GPI[n] (input only) is input.
  - A port pin whose reset default is not GPIO gets it appended: PTA5/~{RESET}.
  - Alternates come from "Composite Pin Name"; directions from the IO Signal Table.
  - NXP's active-low "_b" suffix becomes an overbar (KLC S4.7).
  - Supplies default to power_in; set exceptions from the datasheet with --type NAME=TYPE.
"""
import argparse
import csv
import re
import sys

DIR_TYPE = {"I": "input", "O": "output", "I/O": "bidirectional"}
GROUND = re.compile(r"^(VSS.*|GND.*|VREFL)$")
GPIO = re.compile(r"^GPI?O?\[")


def overbar(name):
    return re.sub(r"^(.*)_b$", r"~{\1}", name)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--device", required=True, help="sheet prefix, e.g. S32K344")
    ap.add_argument("--package", required=True, help="package column suffix, e.g. 257bga")
    ap.add_argument("--type", action="append", default=[], metavar="NAME=TYPE",
                    help="electrical type override for a supply or misc pin (from the datasheet)")
    ap.add_argument("-o", "--out", required=True)
    args = ap.parse_args()
    try:
        import openpyxl
    except ImportError:
        sys.exit("nxp_iomux: openpyxl is required (run scripts/setup_tools.sh, then use its Python)")

    overrides = dict(t.split("=", 1) for t in args.type)
    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    dev, pkg_col = args.device, f"{args.device}_{args.package}"

    ws = wb[f"{dev}_Pinout"]
    hdr = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    col = {h: i for i, h in enumerate(hdr) if h}
    if pkg_col not in col:
        sys.exit(f"nxp_iomux: no column {pkg_col}; have {[h for h in col if h.startswith(dev)]}")
    balls = {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        ball, name = r[col[pkg_col]], r[col["Pin Name"]]
        if ball and name:
            balls[str(ball).strip()] = (str(name).strip(), str(r[col["DEFAULT"]] or "").strip(),
                                        str(r[col["Composite Pin Name"]] or "").strip())

    direction, port = {}, None
    for r in wb[f"{dev}_IO Signal Table"].iter_rows(min_row=3, values_only=True):
        if r[0]:
            port = str(r[0]).strip()
        if port and r[3] and r[6]:
            direction[(port, str(r[3]).strip())] = r[6]

    supplies = set()
    for r in wb[f"{dev}_Supplies"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            supplies.add(str(r[0]).strip())

    rows = []
    power_t = power_b = power_l = power_r = 0
    for ball, (name, default, composite) in sorted(balls.items()):
        m = re.match(r"PT([A-Z])(\d+)$", name)
        if m:
            bit = int(m.group(2))
            label = name
            if default and not GPIO.match(default):
                label += "/" + "/".join(overbar(x) for x in default.split("/"))
            etype = "input" if default.startswith("GPI[") else "bidirectional"
            alts, seen = [], set()
            for f in (x.strip() for x in composite.split("/")):
                if f and not GPIO.match(f) and f not in seen:
                    seen.add(f)
                    alts.append(f"{overbar(f)}:{DIR_TYPE.get(direction.get((name, f)), 'bidirectional')}")
            rows.append([ball, label, etype, f"PT{m.group(1)}", "L" if bit < 16 else "R", bit % 16, ";".join(alts)])
        elif name in supplies:
            etype = overrides.get(name, "power_in")
            side = "B" if GROUND.match(name) else "R" if etype == "power_out" else "T"
            rows.append([ball, name, etype, "Power", side, "", ""])
        else:  # oscillator and other analog misc pins
            etype = overrides.get(name, "input" if name in ("EXTAL", "XTAL32_IN") else
                                  "output" if name in ("XTAL", "XTAL32_OUT") else "passive")
            rows.append([ball, name, etype, "Power", "L", "", ""])

    # Stable, readable order: Power unit first (by name), then ports by bit
    def ball(n):
        m = re.match(r"([A-Z]*)(\d+)$", n)
        return (m.group(1), int(m.group(2))) if m else (n, 0)

    def key(r):
        return (r[3] != "Power", r[3], r[4], (r[1], ball(r[0])) if r[3] == "Power" else (int(r[5]),))
    rows.sort(key=key)
    with open(args.out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["number", "name", "type", "unit", "side", "slot", "alternates"])
        w.writerows(rows)
    ports = sorted({r[3] for r in rows if r[3] != "Power"})
    print(f"{args.out}: {len(rows)} pins; units Power + {', '.join(ports)}")


if __name__ == "__main__":
    main()
